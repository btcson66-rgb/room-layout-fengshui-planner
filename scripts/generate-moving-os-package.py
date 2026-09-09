from __future__ import annotations

import json
import re
import shutil
import textwrap
from datetime import date, timedelta
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "product-output" / "RoomFeng-Moving-New-Home-OS-v1.0"
PDF_OUT = ROOT / "output" / "pdf"
FONT = Path(r"C:\Windows\Fonts\NotoSansTC-VF.ttf")
GREEN = "183C36"
ORANGE = "C76F3C"
PALE = "E8EEE9"
LINE = "AAB7B0"


def reset_dirs() -> None:
    for relative in ["01-Quick-Start", "02-Interactive", "03-Spreadsheets", "04-Printable", "05-Examples", "06-Backups", "07-Preview-Images"]:
        (PACKAGE / relative).mkdir(parents=True, exist_ok=True)
    PDF_OUT.mkdir(parents=True, exist_ok=True)


def write_text_files() -> None:
    (PACKAGE / "README.txt").write_text(
        "RoomFeng Moving & New Home OS v1.0.0\n"
        "====================================\n\n"
        "Start here:\n"
        "1. Read 01-Quick-Start/Quick-Start-Guide.pdf.\n"
        "2. Open the interactive edition using 02-Interactive/ACCESS.txt.\n"
        "3. Use the XLSX workbook when you prefer Excel or Google Sheets.\n"
        "4. Print the A4 or US Letter planner in 04-Printable.\n"
        "5. Import 05-Examples/example-moving-project.json to explore a complete example.\n"
        "6. Export a JSON backup before moving day.\n\n"
        "Online activation: open the stable app URL in ACCESS.txt, choose Payhip or Gumroad, and enter the license key from that marketplace receipt. Activation requires internet; project data remains local to your browser.\n\n"
        "Updates: when a v1.x file update is published, new buyers receive the latest file and existing buyers can re-download it from their Payhip or Gumroad purchase page, subject to the configured download limit. This does not promise lifetime updates.\n\n"
        "Important: Fit results are preliminary measurement checks. Complex 3D rotation, material flexibility, site obstacles, and mover technique are not modelled. Confirm large or expensive items with the mover on site.\n",
        encoding="utf-8",
    )
    (PACKAGE / "LICENSE-OR-USAGE.txt").write_text(
        "PERSONAL USE LICENSE\n\n"
        "The purchaser may use these files personally, print copies for personal use, keep personal backups, and use the licensed online edition on their own devices within the activation policy.\n\n"
        "You may not resell, redistribute, upload the files publicly, share the license key, sublicense, or repackage any part as another product. RoomFeng remains the product author.\n\n"
        "This planning product does not provide architectural, structural, legal, or professional moving advice, and it does not guarantee furniture fit. Refund and purchase terms are supplied by the actual sales platform.\n",
        encoding="utf-8",
    )
    (PACKAGE / "02-Interactive" / "ACCESS.txt").write_text(
        "Interactive edition access\n\n"
        "1. Open the language-specific activation page:\n"
        "- https://roomfeng.win/en/moving-new-home-os/activate/\n"
        "- https://roomfeng.win/zh/moving-new-home-os/activate/\n"
        "2. Choose the marketplace where you purchased (Payhip or Gumroad), then enter that receipt's license key.\n"
        "3. After activation, use the stable app URL:\n"
        "- https://roomfeng.win/en/moving-new-home-os/app/\n"
        "- https://roomfeng.win/zh/moving-new-home-os/app/\n\n"
        "Activation requires internet. The app then supports a short offline grace period and periodically requires re-verification. Project data stays in the current browser unless you export it. Keep the JSON backup in a location you control.\n\n"
        "Purchased but cannot activate? Check the key, check your connection, retry once, then use the support contact shown on the RoomFeng sales page.\n",
        encoding="utf-8",
    )
    (PACKAGE / "06-Backups" / "BACKUP-AND-RECOVERY.txt").write_text(
        "BACKUP\n1. Open Exports & backup.\n2. Choose Export JSON backup.\n3. Store the file in a location you control.\n\nRECOVERY\n1. Open the same module.\n2. Choose Import JSON backup.\n3. Select a RoomFeng JSON backup under 5 MB.\n4. Confirm the dashboard counts and project name.\n\nImported files are validated for schema version, required lists, list size, and overall size before they replace the current local project.\n",
        encoding="utf-8",
    )


def sample_data() -> dict:
    move_date = date.today() + timedelta(days=42)
    return {
        "schemaVersion": 1,
        "locale": "en",
        "exampleData": True,
        "updatedAt": date.today().isoformat(),
        "project": {"name": "Example: Studio to 2-bedroom move", "currentHome": "City studio", "newHome": "2-bedroom apartment", "moveDate": move_date.isoformat(), "currency": "USD", "unit": "cm", "householdSize": 2, "bedrooms": 2, "moveType": "couple"},
        "rooms": [
            {"id": "room-living", "name": "Living room", "lengthMm": 4800, "widthMm": 3600, "ceilingMm": 2800, "doors": [{"id": "door-living", "name": "Entry", "widthMm": 900, "heightMm": 2100, "swing": "left"}], "windows": [{"id": "window-living", "name": "Balcony window", "widthMm": 1800, "heightMm": 1500, "sillHeightMm": 700}], "fixedObjects": [{"id": "column", "type": "column", "note": "40 x 40 cm column, northeast corner"}]},
            {"id": "room-bedroom", "name": "Main bedroom", "lengthMm": 3600, "widthMm": 3200, "ceilingMm": 2800, "doors": [{"id": "door-bedroom", "name": "Bedroom door", "widthMm": 820, "heightMm": 2050, "swing": "right"}], "windows": [], "fixedObjects": []},
        ],
        "entryRoute": {"oldHomeExitWidthMm": 880, "oldHomeExitHeightMm": 2050, "elevatorWidthMm": 1600, "elevatorDepthMm": 1500, "elevatorHeightMm": 2300, "elevatorDoorWidthMm": 900, "corridorWidthMm": 1050, "stairWidthMm": 1000, "landingWidthMm": 1600, "landingDepthMm": 1500, "entranceWidthMm": 920, "entranceHeightMm": 2100, "interiorDoorWidthMm": 820, "interiorDoorHeightMm": 2050},
        "furniture": [], "timeline": [], "boxes": [], "budget": [], "shopping": [], "layouts": [],
        "moveDay": {"moverContact": "Example Moving Co. / 0900-000-000", "buildingContact": "Example building desk", "parking": "Rear loading area 09:00-12:00", "elevatorReservation": "Reserved 09:00-11:00", "notes": "Example data. Replace with real contacts."},
        "firstNight": [],
    }


def populate_sample(data: dict) -> dict:
    categories = [("Queen bed frame", "bed", "room-bedroom", 1520, 2000, 1000), ("Queen mattress", "mattress", "room-bedroom", 1520, 2000, 280), ("Three-seat sofa", "sofa", "room-living", 2100, 900, 850), ("Double wardrobe", "wardrobe", "room-bedroom", 1200, 600, 2200), ("Work desk", "desk", "room-living", 1200, 600, 740), ("Dining table", "dining-table", "room-living", 1400, 800, 750), ("TV console", "tv", "room-living", 1600, 420, 520), ("Bookcase", "shelving", "room-living", 900, 320, 1800)]
    for i, (name, category, room_id, width, depth, height) in enumerate(categories, 1):
        disassemble = category in {"bed", "wardrobe", "desk", "dining-table", "shelving"}
        data["furniture"].append({"id": f"FUR-{i:03d}", "name": name, "category": category, "currentRoom": "Studio", "destinationRoomId": room_id, "dimensions": {"width": width, "depth": depth, "height": height}, "canDisassemble": disassemble, "disassembledDimensions": {"width": min(width, 700), "depth": 180, "height": min(height, 2000)} if disassemble else None, "fragile": category == "tv", "expensive": category == "mattress", "requiresMovers": category in {"mattress", "sofa", "wardrobe"}, "decision": "keep", "notes": ""})
    tasks = ["Book movers", "Measure elevator and turns", "Confirm building move-in rules", "Schedule internet installation", "Sell furniture not moving", "Pack valuables separately", "Prepare first-night box", "Defrost refrigerator", "Label furniture destinations", "Photograph rental condition"]
    for i, title in enumerate(tasks, 1):
        data["timeline"].append({"id": f"TASK-{i:02d}", "title": title, "dueDate": (date.today() + timedelta(days=i * 3)).isoformat(), "phase": "Move plan", "completed": i < 3})
    for i in range(12):
        data["boxes"].append({"id": f"{'LIV' if i < 6 else 'BDR'}-{i+1:03d}", "roomId": "room-living" if i < 6 else "room-bedroom", "contents": ["Books", "Kitchenware", "Winter clothes", "Bedding"][i % 4], "fragile": i == 1, "priority": "high" if i < 2 else "normal", "firstNight": i == 7, "openFirst": i == 7, "packed": i < 7, "loaded": False, "delivered": False, "unpacked": False})
    for i, (cat, label, est, actual) in enumerate([("movers", "Movers", 650, 640), ("packing", "Packing supplies", 95, 82), ("cleaning", "Move-out cleaning", 160, 0), ("internet", "Internet installation", 70, 0)], 1):
        data["budget"].append({"id": f"BUD-{i:02d}", "category": cat, "label": label, "estimated": est, "actual": actual, "paid": i == 2, "dueDate": (date.today() + timedelta(days=14 + i * 5)).isoformat(), "notes": ""})
    for candidate, product, price, depth in [("A", "Compact fabric sofa", 699, 820), ("B", "Two-seat modular sofa", 790, 880), ("C", "Shallow sofa", 620, 760)]:
        data["shopping"].append({"id": f"SHOP-{candidate}", "group": "sofa", "candidate": candidate, "product": product, "category": "sofa", "store": f"Store {candidate}", "url": f"https://example.com/sofa-{candidate.lower()}", "price": price, "dimensions": {"width": 1800, "depth": depth, "height": 790}, "roomId": "room-living", "priority": "normal", "status": "research", "delivery": "7-14 days", "notes": ""})
    data["firstNight"] = [{"id": f"NIGHT-{i:02d}", "label": label, "packed": i < 5} for i, label in enumerate(["Toiletries", "Medication", "Phone charger", "Bedding", "Towel", "Change of clothes", "Water", "Cleaning supplies", "Toilet paper", "Important documents"], 1)]
    return data


def make_workbook(data: dict) -> Path:
    path = PACKAGE / "03-Spreadsheets" / "RoomFeng-Moving-New-Home-OS.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color=LINE)
    input_fill = PatternFill("solid", fgColor="FFF7E8")
    heading_fill = PatternFill("solid", fgColor=GREEN)
    heading_font = Font(color="FFFFFF", bold=True)

    def sheet(name: str, headers: list[str], rows: list[list]) -> object:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}{max(ws.max_row, 2)}"
        for cell in ws[1]:
            cell.fill = heading_fill; cell.font = heading_font; cell.alignment = Alignment(wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.fill = input_fill; cell.border = Border(bottom=thin); cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(1, col).column_letter].width = min(34, max(12, max(len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 60) + 1)) + 2))
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.page_setup.orientation = "landscape" if len(headers) > 7 else "portrait"; ws.page_margins.left = .25; ws.page_margins.right = .25
        ws.print_area = f"A1:{ws.cell(1, len(headers)).column_letter}{max(ws.max_row, 2)}"
        ws.oddFooter.center.text = "RoomFeng Moving & New Home OS | Page &P of &N"
        return ws

    dashboard = wb.create_sheet("Dashboard")
    dashboard.merge_cells("A1:D1"); dashboard["A1"] = "RoomFeng Moving & New Home OS"; dashboard["A1"].fill = heading_fill; dashboard["A1"].font = Font(color="FFFFFF", bold=True, size=20)
    dashboard.append(["Project", data["project"]["name"], "Move date", date.fromisoformat(data["project"]["moveDate"])]); dashboard.append(["Current home", data["project"]["currentHome"], "New home", data["project"]["newHome"]]); dashboard.append([]); dashboard.append(["Metric", "Value", "Metric", "Value"])
    dashboard.append(["Furniture checked", "=COUNTA('Furniture Inventory'!A2:A5000)", "Boxes packed", '=COUNTIF(Boxes!G2:G5000,"Yes")'])
    dashboard.append(["Tasks completed", '=COUNTIF(\'Move Timeline\'!E2:E5000,"Yes")', "Budget estimate", "=SUM(Budget!D2:D5000)"])
    dashboard.append(["Actual spend", "=SUM(Budget!E2:E5000)", "Remaining", "=SUM(Budget!D2:D5000)-SUM(Budget!E2:E5000)"])
    dashboard.freeze_panes = "A5"; dashboard.column_dimensions["A"].width = 24; dashboard.column_dimensions["B"].width = 24; dashboard.column_dimensions["C"].width = 24; dashboard.column_dimensions["D"].width = 24
    for cell in dashboard[5]: cell.fill = heading_fill; cell.font = heading_font
    dashboard.print_area = "A1:D8"; dashboard.sheet_properties.pageSetUpPr.fitToPage = True; dashboard.page_setup.fitToWidth = 1; dashboard.page_setup.fitToHeight = 1

    dashboard["D2"].number_format = "yyyy-mm-dd"
    for coordinate in ["D6", "B7", "D7"]: dashboard[coordinate].number_format = '$#,##0.00;[Red]($#,##0.00);-'
    timeline = sheet("Move Timeline", ["ID", "Phase", "Task", "Due date", "Completed"], [[x["id"], x["phase"], x["title"], date.fromisoformat(x["dueDate"]), "Yes" if x["completed"] else "No"] for x in data["timeline"]])
    furniture = sheet("Furniture Inventory", ["ID", "Name", "Category", "Current room", "Destination", "Width mm", "Depth mm", "Height mm", "Disassemble", "Decision", "Entry", "Room fit", "Usability", "Recommendation", "Notes"], [[x["id"], x["name"], x["category"], x["currentRoom"], x["destinationRoomId"], x["dimensions"]["width"], x["dimensions"]["depth"], x["dimensions"]["height"], "Yes" if x["canDisassemble"] else "No", x["decision"], "Review", "Review", "Review", "Review with mover", x["notes"]] for x in data["furniture"]])
    measurement_rows = [[x["name"], x["lengthMm"], x["widthMm"], x["ceilingMm"], "; ".join(f"{v['name']} {v['widthMm']}x{v['heightMm']}" for v in x["doors"]), "; ".join(v["name"] for v in x["windows"]), "; ".join(v["note"] for v in x["fixedObjects"])] for x in data["rooms"]]
    route = data["entryRoute"]
    measurement_rows.append(["Entry route", "", "", "", f"Old exit {route['oldHomeExitWidthMm']}x{route['oldHomeExitHeightMm']}; elevator door {route['elevatorDoorWidthMm']}; entrance {route['entranceWidthMm']}x{route['entranceHeightMm']}; interior door {route['interiorDoorWidthMm']}x{route['interiorDoorHeightMm']}", "", f"Elevator {route['elevatorWidthMm']}x{route['elevatorDepthMm']}x{route['elevatorHeightMm']}; corridor {route['corridorWidthMm']}; stair {route['stairWidthMm']}; landing {route['landingWidthMm']}x{route['landingDepthMm']}"])
    measurements = sheet("Measurements", ["Room / route", "Length mm", "Width mm", "Ceiling mm", "Doors", "Windows", "Fixed objects / route details"], measurement_rows)
    shopping = sheet("Furniture Shopping", ["Group", "Candidate", "Product", "Category", "Store", "URL", "Price", "Width mm", "Depth mm", "Height mm", "Room", "Fit", "Status", "Delivery", "Rating", "Notes"], [[x["group"], x["candidate"], x["product"], x["category"], x["store"], x["url"], x["price"], x["dimensions"]["width"], x["dimensions"]["depth"], x["dimensions"]["height"], x["roomId"], "Layout review", x["status"], x["delivery"], "", x["notes"]] for x in data["shopping"]])
    boxes = sheet("Boxes", ["Box ID", "Destination", "Contents", "Fragile", "Priority", "First night", "Packed", "Loaded", "Delivered", "Unpacked"], [[x["id"], x["roomId"], x["contents"], "Yes" if x["fragile"] else "No", x["priority"], "Yes" if x["firstNight"] else "No", "Yes" if x["packed"] else "No", "No", "No", "No"] for x in data["boxes"]])
    budget = sheet("Budget", ["ID", "Category", "Item", "Estimated", "Actual", "Paid", "Due date", "Notes"], [[x["id"], x["category"], x["label"], x["estimated"], x["actual"], "Yes" if x["paid"] else "No", date.fromisoformat(x["dueDate"]) if x["dueDate"] else None, x["notes"]] for x in data["budget"]])
    move_day_rows = [["New home", data["project"]["newHome"]], ["Mover contact", data["moveDay"]["moverContact"]], ["Building contact", data["moveDay"]["buildingContact"]], ["Parking", data["moveDay"]["parking"]], ["Elevator reservation", data["moveDay"]["elevatorReservation"]]]
    for index, room in enumerate(data["rooms"], start=1):
        furniture_names = ", ".join(item["name"] for item in data["furniture"] if item["destinationRoomId"] == room["id"])
        move_day_rows.append([f"Room {index:02d} {room['name']}", furniture_names or "—"])
    move_day_rows.extend([["Sequence", "1. Large furniture  2. Furthest room  3. Smaller furniture  4. Boxes"], ["Notes", data["moveDay"]["notes"]]])
    move_day = sheet("Move Day", ["Section", "Details"], move_day_rows)
    instructions = sheet("Instructions", ["Step", "Action"], [[1, "Set up the project and move date."], [2, "Measure the full entry route and each room."], [3, "Add the largest furniture first and review all three gates."], [4, "Plan tasks, boxes, budget, and shopping candidates."], [5, "Export a JSON backup before moving day."], [6, "Preliminary measurement check only; confirm critical items on site."]])

    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    furniture.add_data_validation(yes_no); yes_no.add("I2:I5000")
    for ws, ranges in [(timeline, ["E2:E5000"]), (boxes, ["D2:D5000", "F2:J5000"]), (budget, ["F2:F5000"])]:
        dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True); ws.add_data_validation(dv)
        for target in ranges: dv.add(target)
    decision = DataValidation(type="list", formula1='"Keep,Sell,Donate,Dispose,Undecided"', allow_blank=True); furniture.add_data_validation(decision); decision.add("J2:J5000")
    candidate = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True); shopping.add_data_validation(candidate); candidate.add("B2:B5000")
    budget["D2"].number_format = '$#,##0.00;[Red]($#,##0.00);-'
    for row in budget.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row: cell.number_format = '$#,##0.00;[Red]($#,##0.00);-'
    for row in timeline.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row: cell.number_format = "yyyy-mm-dd"
    for row in budget.iter_rows(min_row=2, min_col=7, max_col=7):
        for cell in row: cell.number_format = "yyyy-mm-dd"
    red_fill = PatternFill("solid", fgColor="F7D9D6")
    furniture.conditional_formatting.add("K2:M5000", FormulaRule(formula=['OR(K2="Fail",L2="Fail",M2="Conflict")'], fill=red_fill))
    wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.calculation.calcMode = "auto"
    wb.save(path)
    cached_values = {
        "B5": len(data["furniture"]),
        "D5": sum(1 for box in data["boxes"] if box["packed"]),
        "B6": sum(1 for task in data["timeline"] if task["completed"]),
        "D6": sum(item["estimated"] for item in data["budget"]),
        "B7": sum(item["actual"] for item in data["budget"]),
    }
    cached_values["D7"] = cached_values["D6"] - cached_values["B7"]
    cache_workbook_formula_results(path, cached_values)
    return path


def cache_workbook_formula_results(path: Path, values: dict[str, int | float]) -> None:
    """Keep formula cells useful in viewers that do not run Excel's calculation engine."""
    replacement = path.with_suffix(".cached.xlsx")
    with ZipFile(path, "r") as source, ZipFile(replacement, "w", ZIP_DEFLATED) as destination:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = content.decode("utf-8")
                for coordinate, value in values.items():
                    pattern = rf'(<c r="{coordinate}"[^>]*>.*?<f[^>]*>.*?</f><v>)(.*?)(</v>)'
                    xml, count = re.subn(pattern, rf'\g<1>{value}\g<3>', xml, count=1, flags=re.DOTALL)
                    if count != 1:
                        raise RuntimeError(f"Unable to cache Dashboard formula result for {coordinate}")
                content = xml.encode("utf-8")
            destination.writestr(info, content)
    replacement.replace(path)


def styles():
    pdfmetrics.registerFont(TTFont("NotoTC", str(FONT)))
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("TitleTC", parent=base["Title"], fontName="NotoTC", fontSize=25, leading=31, textColor=colors.HexColor("#183C36"), alignment=TA_CENTER, spaceAfter=18),
        "h1": ParagraphStyle("H1TC", parent=base["Heading1"], fontName="NotoTC", fontSize=18, leading=23, textColor=colors.HexColor("#183C36"), spaceAfter=10),
        "h2": ParagraphStyle("H2TC", parent=base["Heading2"], fontName="NotoTC", fontSize=13, leading=18, textColor=colors.HexColor("#C76F3C"), spaceAfter=7),
        "body": ParagraphStyle("BodyTC", parent=base["BodyText"], fontName="NotoTC", fontSize=9.5, leading=15, textColor=colors.HexColor("#263B35"), spaceAfter=7),
        "small": ParagraphStyle("SmallTC", parent=base["BodyText"], fontName="NotoTC", fontSize=7.5, leading=11, textColor=colors.HexColor("#53665F")),
    }


def footer(canvas, doc):
    canvas.saveState(); canvas.setFont("NotoTC", 7); canvas.setFillColor(colors.HexColor("#60726B")); canvas.drawString(doc.leftMargin, 9 * mm, "RoomFeng Moving & New Home OS - preliminary planning tool"); canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 9 * mm, f"{canvas.getPageNumber()}"); canvas.restoreState()


def planner_pdf(path: Path, page_size, lang: str = "en", quick: bool = False) -> None:
    s = styles(); zh = lang == "zh"
    doc = SimpleDocTemplate(str(path), pagesize=page_size, rightMargin=15*mm, leftMargin=15*mm, topMargin=17*mm, bottomMargin=16*mm, title="RoomFeng Moving & New Home OS")
    story = [Spacer(1, 18*mm), Paragraph("RoomFeng", s["h2"]), Paragraph("搬家與新居規劃系統" if zh else "Moving & New Home OS", s["title"]), Paragraph("MOVE - MEASURE - FIT - PLAN - BUY - PACK - MOVE IN", s["small"]), Spacer(1, 8*mm), Paragraph("這是一套減少錯誤、遺漏與重工的搬家工作系統。" if zh else "A working system for fewer mistakes, fewer missed tasks, and less rework.", s["body"]), Spacer(1, 20*mm), Paragraph("重要：本工具為搬家前尺寸初步檢查，不取代搬家公司現場評估。" if zh else "Important: This is a preliminary measurement check and does not replace an on-site mover assessment.", s["body"]), PageBreak()]
    steps = [("1", "建立專案" if zh else "Set up the project", "日期、住家、幣別與單位。" if zh else "Move date, homes, currency, and unit."), ("2", "量完整路線" if zh else "Measure the full route", "從舊家出口到新家目的房間。" if zh else "From the old-home exit to the destination room."), ("3", "先加最大家具" if zh else "Add the largest furniture", "查看搬得進去、放得下、用得順。" if zh else "Review entry, room fit, and usability."), ("4", "排工作與紙箱" if zh else "Plan tasks and boxes", "讓 Dashboard 告訴你下一步。" if zh else "Let the dashboard identify the next action."), ("5", "匯出與備份" if zh else "Export and back up", "搬家日前保存 JSON 與列印指揮表。" if zh else "Save JSON and print the command sheet before move day.")]
    story += [Paragraph("5 分鐘快速開始" if zh else "Five-minute quick start", s["h1"]), Table([[Paragraph(n,s["h2"]), Paragraph(f"<b>{title}</b><br/>{copy}",s["body"])] for n,title,copy in steps], colWidths=[15*mm,145*mm], style=TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LINEBELOW",(0,0),(-1,-2),.4,colors.HexColor("#C3CCC6")),("BOTTOMPADDING",(0,0),(-1,-1),8)]))]
    if quick:
        story += [Spacer(1, 5*mm), Paragraph("已購買但無法啟用？" if zh else "Purchased but can't activate?", s["h2"]), Paragraph("1. 檢查授權碼與多餘空格。 2. 確認網路。 3. 稍後重試。 4. 使用銷售頁上的正式支援聯絡方式。" if zh else "1. Check the license key and extra spaces. 2. Check the internet connection. 3. Retry after a short wait. 4. Use the production support contact shown on the sales page.", s["body"]), Paragraph("授權驗證需要網路；搬家專案資料仍留在瀏覽器。" if zh else "License verification needs internet; moving project data remains in the browser.", s["small"])]
        doc.build(story, onFirstPage=footer, onLaterPages=footer); return
    sections = [
        ("Home measurements", ["Room: __________________________", "Length: __________  Width: __________  Ceiling: __________", "Door width / height: __________________  Swing: __________", "Window width / height / sill: ______________________________", "Fixed objects and obstructions: __________________________________________"]),
        ("Entry route", ["Old-home exit: W ______ H ______", "Elevator inside: W ______ D ______ H ______", "Elevator door: ______  Corridor: ______  Stair: ______", "Landing: W ______ D ______", "New-home entrance: W ______ H ______", "Interior doorway: W ______ H ______"]),
        ("Furniture inventory", ["ID", "Furniture", "W x D x H", "Room", "Keep / Sell", "Three-gate result"]),
        ("Furniture decision", ["[ ] Likely fit  [ ] Review required  [ ] Current measurements fail", "Why: _________________________________________________________________", "[ ] Keep  [ ] Disassemble first  [ ] Review with mover  [ ] Sell / replace"]),
        ("Furniture shopping A / B / C", ["Candidate", "Price", "W x D x H", "Fit", "Delivery", "Notes"]),
        ("Moving timeline", ["8 weeks", "6 weeks", "4 weeks", "2 weeks", "1 week", "3 days", "Move day", "First night", "First week"]),
        ("Box tracker", ["Box ID", "Destination", "Contents", "Fragile", "Packed", "Loaded", "Delivered", "Unpacked"]),
        ("Budget", ["Category / item", "Estimated", "Actual", "Paid", "Due", "Notes"]),
        ("Move-day command sheet", ["Mover contact: ______________________________", "Building contact: ___________________________", "Parking / loading: __________________________", "Elevator reservation: _______________________", "Sequence: 1. Large furniture  2. Furthest room  3. Smaller furniture  4. Boxes"]),
        ("First-night kit", ["[ ] Toiletries", "[ ] Medication", "[ ] Phone charger", "[ ] Bedding", "[ ] Towel", "[ ] Change of clothes", "[ ] Water", "[ ] Cleaning supplies", "[ ] Toilet paper", "[ ] Important documents"]),
        ("Notes", ["", "", "", "", "", "", "", ""]),
    ]
    for title, rows in sections:
        story += [PageBreak(), Paragraph(title, s["h1"])]
        if len(rows) and rows[0] in {"ID", "Candidate", "Box ID", "Category / item"}:
            table_rows = [rows] + [["" for _ in rows] for _ in range(10)]
            widths = [doc.width / len(rows)] * len(rows)
            table = Table(table_rows, colWidths=widths, rowHeights=[8*mm] + [12*mm]*10, repeatRows=1)
            table.setStyle(TableStyle([("FONT",(0,0),(-1,-1),"NotoTC",7),("BACKGROUND",(0,0),(-1,0),colors.HexColor("#183C36")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.45,colors.HexColor("#9AABA3")),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),4)]))
            story.append(table)
        else:
            story.append(Table([[Paragraph(row or "&nbsp;", s["body"])] for row in rows], colWidths=[doc.width], rowHeights=[13*mm]*len(rows), style=TableStyle([("LINEBELOW",(0,0),(-1,-1),.45,colors.HexColor("#AAB7B0")),("VALIGN",(0,0),(-1,-1),"BOTTOM")])))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def draw_preview(path: Path, number: int, title: str, subtitle: str, mode: str) -> None:
    image = Image.new("RGB", (1600, 1000), "#E9ECE6"); d = ImageDraw.Draw(image)
    font_title = ImageFont.truetype(str(FONT), 62); font_sub = ImageFont.truetype(str(FONT), 28); font_small = ImageFont.truetype(str(FONT), 22)
    d.rectangle((0,0,1600,155), fill="#183C36"); d.text((72,48), "ROOMFENG / MOVING OS", font=font_sub, fill="#F5F2E9"); d.text((1390,32), f"{number:02d}", font=font_title, fill="#E59A6A")
    d.text((72,205), title, font=font_title, fill="#183C36"); d.text((75,290), subtitle, font=font_sub, fill="#556A63")
    if mode == "hero":
        d.rectangle((75,380,820,895), fill="#FBFAF5", outline="#183C36", width=4)
        d.line((165,500,710,500,710,790,165,790,165,500), fill="#183C36", width=12)
        d.arc((165,630,345,810), 180, 270, fill="#C76F3C", width=8)
        d.rectangle((455,565,650,690), fill="#DDE8E1", outline="#183C36", width=5)
        d.text((115,825), "ROOM + ROUTE PLANNER", font=font_small, fill="#556A63")
        hero_cards = [("FURNITURE FIT", "LIKELY FIT"), ("BUDGET", "$263 LEFT"), ("PACKING", "21 / 42 BOXES"), ("NEXT STEP", "MOVE DAY")]
        for i, (label, value) in enumerate(hero_cards):
            x = 880 + (i % 2) * 340; y = 380 + (i // 2) * 260
            d.rectangle((x, y, x + 300, y + 220), fill="#FBFAF5", outline="#AAB7B0", width=3)
            d.text((x + 24, y + 28), label, font=font_small, fill="#C76F3C")
            d.text((x + 24, y + 105), value, font=font_sub, fill="#183C36")
    elif mode == "dashboard":
        values = [("68%", "READY"), ("12 / 18", "FURNITURE"), ("21 / 42", "BOXES"), ("31 / 44", "TASKS")]
        for i,(value,label) in enumerate(values):
            x=75+(i%2)*620; y=390+(i//2)*230; d.rectangle((x,y,x+560,y+180),fill="#FBFAF5",outline="#AAB7B0",width=3); d.text((x+35,y+25),value,font=font_title,fill="#183C36"); d.text((x+38,y+112),label,font=font_small,fill="#C76F3C")
        d.rectangle((1325,390,1525,800),fill="#183C36"); d.rectangle((1375,500,1475,760),fill="#E59A6A")
    elif mode == "budget":
        totals = [("ESTIMATED", "$1,245"), ("ACTUAL", "$982"), ("REMAINING", "$263")]
        for i, (label, value) in enumerate(totals):
            x = 75 + i * 495
            d.rectangle((x, 390, x + 445, 555), fill="#FBFAF5", outline="#AAB7B0", width=3)
            d.text((x + 28, 420), label, font=font_small, fill="#C76F3C")
            d.text((x + 28, 465), value, font=font_title, fill="#183C36")
        categories = [("MOVERS", 690), ("PACKING", 190), ("CLEANING", 150), ("INTERNET", 75)]
        for i, (label, value) in enumerate(categories):
            y = 625 + i * 62
            d.text((75, y), label, font=font_small, fill="#556A63")
            d.rectangle((300, y + 4, 300 + value, y + 30), fill="#183C36" if i == 0 else "#C76F3C")
    elif mode == "fit":
        for i,(gate,status,color) in enumerate([("GATE 1","LIKELY FIT","#2C7258"),("GATE 2","PASS","#2C7258"),("GATE 3","REVIEW","#A46C1F")]):
            y=390+i*155; d.rectangle((75,y,1170,y+120),fill="#FBFAF5",outline=color,width=5); d.text((110,y+26),gate,font=font_small,fill="#566B64"); d.text((380,y+20),status,font=font_title,fill=color); d.ellipse((1250,y+10,1360,y+120),fill=color); d.text((1285,y+30),"!" if status=="REVIEW" else "✓",font=font_title,fill="white")
    elif mode == "timeline":
        d.line((160,420,160,890),fill="#183C36",width=8)
        for i,label in enumerate(["8 WEEKS","4 WEEKS","1 WEEK","MOVE DAY"]): y=430+i*135; d.ellipse((135,y,185,y+50),fill="#C76F3C"); d.text((230,y),label,font=font_sub,fill="#183C36"); d.rectangle((510,y-5,1360,y+65),fill="#FBFAF5",outline="#B8C2BB",width=2)
    elif mode == "boxes":
        for i in range(6):
            x=75+(i%3)*490; y=390+(i//3)*245; d.rectangle((x,y,x+430,y+200),fill="#FBFAF5",outline="#183C36",width=5); d.text((x+25,y+20),f"BDR-{i+1:03d}",font=font_title,fill="#183C36"); d.text((x+28,y+110),"BEDROOM / OPEN FIRST" if i==0 else "BEDROOM",font=font_small,fill="#C76F3C")
    elif mode == "spreadsheet":
        d.rectangle((75,380,1500,900),fill="white",outline="#183C36",width=5); d.rectangle((75,380,1500,450),fill="#183C36");
        for x in [250,520,800,1080,1300]: d.line((x,380,x,900),fill="#BAC4BE",width=2)
        for y in range(450,901,65): d.line((75,y,1500,y),fill="#BAC4BE",width=2)
        d.text((105,395),"FURNITURE INVENTORY",font=font_small,fill="white")
    elif mode == "pdf":
        for i in range(3): x=170+i*390; d.rectangle((x,380,x+310,850),fill="white",outline="#9AA8A0",width=3); d.rectangle((x,380,x+310,450),fill="#183C36"); d.text((x+25,400),["MEASURE","PLAN","MOVE"][i],font=font_small,fill="white"); [d.line((x+30,y,x+280,y),fill="#C6CEC8",width=2) for y in range(500,810,48)]
    elif mode == "move-day":
        d.rectangle((75,370,1500,900),fill="#FBFAF5",outline="#183C36",width=6); d.text((115,410),"MOVE-DAY COMMAND SHEET",font=font_title,fill="#183C36"); d.line((115,500,1450,500),fill="#C76F3C",width=5); d.text((115,550),"01  LIVING ROOM",font=font_sub,fill="#183C36"); d.text((115,620),"02  BEDROOM",font=font_sub,fill="#183C36"); d.text((820,550),"1. LARGE FURNITURE",font=font_sub,fill="#183C36"); d.text((820,620),"2. FURTHEST ROOM",font=font_sub,fill="#183C36")
    else:
        labels=["INTERACTIVE","XLSX","A4 + LETTER","JSON BACKUP"]
        for i,label in enumerate(labels): x=75+(i%2)*730;y=390+(i//2)*220;d.rectangle((x,y,x+660,y+170),fill="#FBFAF5",outline="#183C36",width=4);d.text((x+35,y+45),label,font=font_title,fill="#183C36")
    image.save(path, quality=94)


def generate_previews() -> None:
    specs = [("Plan the move before moving day.","Planner · Furniture fit · Budget · Packing","hero"),("Dashboard","Know what happens next.","dashboard"),("Furniture Fit","Every status includes a reason.","fit"),("Moving Timeline","Count back from move day.","timeline"),("Box Tracker","Find the first-night box fast.","boxes"),("Spreadsheet","Nine practical Excel sheets.","spreadsheet"),("Printable PDF","Designed for paper, not screenshots.","pdf"),("Budget","Estimate, actual, and remaining.","budget"),("Move-Day Sheet","One page for the door and phone.","move-day"),("What's Included","Interactive, spreadsheet, print, backup.","included")]
    for i,(title,subtitle,mode) in enumerate(specs,1): draw_preview(PACKAGE/"07-Preview-Images"/f"preview-{i:02d}-{mode}.png",i,title,subtitle,mode)


def verify_workbook(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    expected = ["Dashboard", "Move Timeline", "Furniture Inventory", "Measurements", "Furniture Shopping", "Boxes", "Budget", "Move Day", "Instructions"]
    assert wb.sheetnames == expected
    for ws in wb.worksheets:
        assert ws.freeze_panes or ws.title == "Instructions"
        assert ws.print_area
    formulas = [cell.value for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f"]
    assert formulas and not any("#REF!" in str(value) or "#VALUE!" in str(value) for value in formulas)


def one_page_extract(source: Path, destination: Path, page_index: int) -> None:
    reader = PdfReader(source); writer = PdfWriter(); writer.add_page(reader.pages[page_index])
    with destination.open("wb") as handle: writer.write(handle)


def generate_free_preview() -> Path:
    preview_dir = ROOT / "output" / "free-preview"; preview_dir.mkdir(parents=True, exist_ok=True)
    quick_extract = preview_dir / "Quick-Start-Extract.pdf"
    planner_sample = preview_dir / "Planner-Sample-Page.pdf"
    one_page_extract(PACKAGE / "01-Quick-Start" / "Quick-Start-Guide.pdf", quick_extract, 1)
    one_page_extract(PACKAGE / "04-Printable" / "RoomFeng-Moving-Planner-A4.pdf", planner_sample, 2)
    destination = ROOT / "public" / "downloads" / "roomfeng-moving-os-free-preview-v1.0.zip"
    destination.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(destination, "w", ZIP_DEFLATED) as archive:
        archive.writestr("README.txt", "RoomFeng Moving OS free preview\n\nIncludes two product images, a one-page quick-start extract, and one planner sample page. This is not the complete paid product.\n")
        archive.write(PACKAGE / "07-Preview-Images" / "preview-02-dashboard.png", "Dashboard-Preview.png")
        archive.write(PACKAGE / "07-Preview-Images" / "preview-03-fit.png", "Furniture-Fit-Preview.png")
        archive.write(quick_extract, quick_extract.name)
        archive.write(planner_sample, planner_sample.name)
    return destination


def main() -> None:
    reset_dirs(); write_text_files(); data = populate_sample(sample_data())
    sample_path = PACKAGE / "05-Examples" / "example-moving-project.json"; sample_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    shutil.copy2(sample_path, PACKAGE / "06-Backups" / "example-backup.json")
    workbook = make_workbook(data); verify_workbook(workbook)
    planner_pdf(PACKAGE / "01-Quick-Start" / "Quick-Start-Guide.pdf", A4, "en", quick=True)
    planner_pdf(PACKAGE / "01-Quick-Start" / "快速開始指南.pdf", A4, "zh", quick=True)
    planner_pdf(PACKAGE / "04-Printable" / "RoomFeng-Moving-Planner-A4.pdf", A4)
    planner_pdf(PACKAGE / "04-Printable" / "RoomFeng-Moving-Planner-US-Letter.pdf", LETTER)
    for pdf in (PACKAGE / "04-Printable").glob("*.pdf"): shutil.copy2(pdf, PDF_OUT / pdf.name)
    generate_previews()
    metadata = {"product": "RoomFeng Moving & New Home OS", "productVersion": "1.0.0", "schemaVersion": 1, "releaseDate": None, "providers": ["Payhip", "Gumroad"], "licenseRequiredForOnlineApp": True}
    (PACKAGE / "PRODUCT-METADATA.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    manifest = {"product": "RoomFeng Moving & New Home OS", "version": "1.0.0", "schemaVersion": 1, "generated": date.today().isoformat(), "files": sorted(str(path.relative_to(PACKAGE)).replace("\\", "/") for path in PACKAGE.rglob("*") if path.is_file() and path.name != "MANIFEST.json")}
    (PACKAGE / "MANIFEST.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    archive = shutil.make_archive(str(PACKAGE), "zip", PACKAGE.parent, PACKAGE.name)
    preview = generate_free_preview()
    print(json.dumps({"package": str(PACKAGE), "archive": archive, "preview": str(preview), "files": len(manifest["files"])}, indent=2))


if __name__ == "__main__":
    main()
