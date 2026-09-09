from pathlib import Path
import json
import re
import tempfile
import zipfile

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "product-output" / "RoomFeng-Moving-New-Home-OS-v1.0"


def verify_xlsx():
    path = PACKAGE / "03-Spreadsheets" / "RoomFeng-Moving-New-Home-OS.xlsx"
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    expected = ["Dashboard", "Move Timeline", "Furniture Inventory", "Measurements", "Furniture Shopping", "Boxes", "Budget", "Move Day", "Instructions"]
    assert formulas.sheetnames == expected
    assert formulas["Dashboard"]["B6"].data_type == "f"
    assert isinstance(values["Dashboard"]["B6"].value, (int, float))
    for ws in formulas.worksheets:
        assert ws.print_area, f"missing print area: {ws.title}"
        assert ws.freeze_panes, f"missing freeze pane: {ws.title}"
        if ws.title != "Dashboard": assert ws.auto_filter.ref, f"missing filter: {ws.title}"
        assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert sum(len(ws.data_validations.dataValidation) for ws in formulas.worksheets) >= 5
    assert sum(len(ws.conditional_formatting) for ws in formulas.worksheets) >= 1
    assert formulas["Budget"]["D2"].number_format != "General"
    assert formulas["Budget"]["G2"].number_format.replace("\\", "") == "yyyy-mm-dd"
    assert formulas["Move Timeline"]["D2"].number_format.replace("\\", "") == "yyyy-mm-dd"
    errors = []
    for ws in values.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A")):
                    errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    assert not errors, errors
    return {"sheets": len(expected), "formula_errors": 0, "dashboard_cached_value": values["Dashboard"]["B6"].value}


def verify_pdfs():
    result = {}
    for path in sorted(PACKAGE.rglob("*.pdf")):
        reader = PdfReader(path)
        assert len(reader.pages) > 0
        text = "".join((page.extract_text() or "") for page in reader.pages)
        assert "RoomFeng" in text
        result[path.name] = {"pages": len(reader.pages), "bytes": path.stat().st_size}
    assert result["RoomFeng-Moving-Planner-A4.pdf"]["pages"] == result["RoomFeng-Moving-Planner-US-Letter.pdf"]["pages"]
    return result


def verify_previews():
    previews = sorted((PACKAGE / "07-Preview-Images").glob("*.png"))
    assert len(previews) == 10
    for path in previews:
        with Image.open(path) as image:
            assert image.size == (1600, 1000)
            assert image.mode == "RGB"
    return {"count": len(previews), "dimensions": "1600x1000"}


def verify_package():
    manifest = json.loads((PACKAGE / "MANIFEST.json").read_text(encoding="utf-8"))
    disk = sorted(str(path.relative_to(PACKAGE)).replace("\\", "/") for path in PACKAGE.rglob("*") if path.is_file() and path.name != "MANIFEST.json")
    assert manifest["files"] == disk
    archive = Path(f"{PACKAGE}.zip")
    assert archive.exists()
    with zipfile.ZipFile(archive) as zf:
        assert zf.testzip() is None
        with tempfile.TemporaryDirectory(prefix="roomfeng-product-qa-") as temporary:
            extracted = Path(temporary)
            zf.extractall(extracted)
            archive_root = extracted / PACKAGE.name
            assert archive_root.is_dir(), "archive must unpack into one named product directory"
            extracted_files = sorted(str(path.relative_to(archive_root)).replace("\\", "/") for path in archive_root.rglob("*") if path.is_file() and path.name != "MANIFEST.json")
            assert extracted_files == disk

    banned = re.compile(r"(?:[A-Za-z]:[\\/](?![\\/])|localhost|127\.0\.0\.1|Fable company|Company Vault|codex[\\/]|\bapi_key\b|product-secret-key|PAYHIP_SECRET|license secret|developer note|source map|debug data|(?:^|[\\/])temp(?:[\\/]|$))", re.IGNORECASE)
    findings = []
    for path in PACKAGE.rglob("*"):
        if not path.is_file(): continue
        text = ""
        if path.suffix.lower() in {".txt", ".json", ".md", ".csv"}: text = path.read_text(encoding="utf-8", errors="replace")
        elif path.suffix.lower() == ".pdf": text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        elif path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, data_only=False, read_only=True)
            text = "\n".join(str(cell.value) for ws in workbook.worksheets for row in ws.iter_rows() for cell in row if cell.value is not None)
        if banned.search(text) or banned.search(str(path.relative_to(PACKAGE))): findings.append(str(path.relative_to(PACKAGE)))
    assert not findings, f"development-machine or secret markers found: {findings}"
    metadata = json.loads((PACKAGE / "PRODUCT-METADATA.json").read_text(encoding="utf-8"))
    assert metadata["productVersion"] == "1.0.0" and metadata["schemaVersion"] == 1 and metadata["releaseDate"] is None
    assert metadata["providers"] == ["Payhip", "Gumroad"]
    assert "PERSONAL USE LICENSE" in (PACKAGE / "LICENSE-OR-USAGE.txt").read_text(encoding="utf-8")
    access = (PACKAGE / "02-Interactive" / "ACCESS.txt").read_text(encoding="utf-8")
    assert "https://roomfeng.win/en/moving-new-home-os/activate/" in access and "Payhip or Gumroad" in access
    return {"manifest_files": len(disk), "zip_bytes": archive.stat().st_size, "extracted_match": True, "sensitive_findings": 0}


def verify_free_preview():
    path = ROOT / "public" / "downloads" / "roomfeng-moving-os-free-preview-v1.0.zip"
    assert path.exists()
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        names = sorted(archive.namelist())
        assert names == ["Dashboard-Preview.png", "Furniture-Fit-Preview.png", "Planner-Sample-Page.pdf", "Quick-Start-Extract.pdf", "README.txt"]
        assert not any(name.endswith((".xlsx", ".json")) for name in names)
    return {"files": 5, "bytes": path.stat().st_size, "contains_paid_core": False}


if __name__ == "__main__":
    print(json.dumps({"xlsx": verify_xlsx(), "pdf": verify_pdfs(), "previews": verify_previews(), "package": verify_package(), "free_preview": verify_free_preview()}, ensure_ascii=False, indent=2))
